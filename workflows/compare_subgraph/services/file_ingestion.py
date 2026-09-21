import io
import openpyxl
import pymupdf
import pandas as pd

def ingest_document(file_ext: str, file_bytes: bytes) -> dict:
    """
    Universal in-memory ingestion engine.
    """
    if not file_bytes:
        return {"content": "", "is_scanned": False, "error": "0-byte file"}

    ext = file_ext.lower()

    try:
        if ext == ".txt":
            return {"content": file_bytes.decode("utf-8"), "is_scanned": False}
        elif ext == ".pdf":
            return parse_pdf_bytes(file_bytes)
        elif ext == ".docx":
            return parse_docx_bytes(file_bytes)
        elif ext == ".xlsx":
            return parse_excel_bytes(file_bytes)
        else:
            return {"content": "", "is_scanned": False, "error": f"Unsupported format: {ext}"}
    except Exception as e:
        return {"content": "", "is_scanned": False, "error": str(e)}

def parse_pdf_bytes(file_bytes: bytes) -> dict:
    # PyMuPDF reads directly from the byte stream
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    markdown_content = []
    total_text_len = 0
    has_images = False

    for page in doc:
        if page.get_images():
            has_images = True

        tabs = page.find_tables()
        if tabs.tables:
            for tab in tabs:
                markdown_content.append(tab.to_markdown())

        text = page.get_text().strip()
        total_text_len += len(text)
        markdown_content.append(text)

    doc.close()

    is_scanned = (total_text_len < 50) and has_images

    return {
        "content": "\n\n".join(markdown_content) if not is_scanned else "",
        "is_scanned": is_scanned
    }

def parse_docx_bytes(file_bytes: bytes) -> dict:
    # Wrap bytes in a file-like object for python-docx
    try:
        import docx
    except ModuleNotFoundError:
        return {
            "content": "",
            "is_scanned": False,
            "error": "python-docx is not installed",
        }
    doc = docx.Document(io.BytesIO(file_bytes))
    markdown_content = []

    for para in doc.paragraphs:
        if para.text.strip():
            markdown_content.append(para.text.strip())

    markdown_content.append("\n### Document Tables\n")

    for table in doc.tables:
        for i, row in enumerate(table.rows):
            row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            markdown_content.append("| " + " | ".join(row_data) + " |")
            if i == 0:
                markdown_content.append("|" + "|".join(["---"] * len(row.cells)) + "|")
        markdown_content.append("\n")

    return {"content": "\n".join(markdown_content), "is_scanned": False}


def parse_excel_bytes(file_bytes: bytes) -> dict:
    # Pandas can read directly from the BytesIO stream
    sheets = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None, engine='openpyxl')
    markdown_content = []

    for sheet_name, df in sheets.items():
        markdown_content.append(f"### Sheet: {sheet_name}\n")
        markdown_content.append(df.fillna("").to_markdown(index=False))
        markdown_content.append("\n")

    return {"content": "\n".join(markdown_content), "is_scanned": False}

